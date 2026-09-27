from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
import json
root=Path(r'C:\Users\rober\OneDrive\Desktop\mit\Information Systyem')
out=Path(r'C:\Users\rober\sol-pump-radar\output\pbl_combined')
selection=[
('MN522_PBL Week1 (1).docx',[35,39,42,46]),
('MN22 pbl week 2 .docx',[32,36,40,44]),
('MN522 pbl week 3 .docx',[15,19,22,25]),
('MN522 PBl week_04.docx',[25,27,29,31]),
('MN522_pbl week 5.docx',[28,30,32,35]),
('MN522_pbl week 6.docx',[24,27,30]),
('MN522_pbl_week 07.docx',[29,33,37,41]),
('MN522_pbl_ week 8.docx',[19,23,27,31]),
('MN522_pbl_week9.docx',[28,32,36,40]),
('MN522_pbl_week10.docx',[25,28,32,36]),
('MN522 pbl week 11.docx',[39,41,43,45])]
d=Document()
s=d.sections[0]
s.page_width=Inches(8.27); s.page_height=Inches(11.69)
s.top_margin=s.bottom_margin=Inches(.8)
s.left_margin=s.right_margin=Inches(.85)
normal=d.styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(12)
normal.paragraph_format.line_spacing=1.08
h=d.styles['Heading 1']; h.font.name='Calibri'; h.font.size=Pt(17); h.font.color.rgb=RGBColor(0,0,0)
h.paragraph_format.space_after=Pt(18)
expected=[]; audit=[]
for week,(filename,indices) in enumerate(selection,1):
 src=Document(root/filename)
 title=d.add_paragraph(f'Week {week} Answers',style='Heading 1')
 title.paragraph_format.page_break_before=week>1
 expected.append(f'Week {week} Answers')
 for number,i in enumerate(indices,1):
  from docx.text.paragraph import Paragraph
  e=list(src.element.body)[i]
  assert e.tag.endswith('}p')
  p=Paragraph(e,src._body)
  assert p.text.strip()
  elements=list(src.element.body)
  qi=i-1
  while qi>=0:
   q=Paragraph(elements[qi],src._body)
   if q.text.strip(): break
   qi-=1
  assert qi>=0
  question=f'{number}. {q.text}'
  qp=d.add_paragraph()
  qp.add_run(question).bold=True
  qp.paragraph_format.keep_with_next=True
  qp.paragraph_format.space_after=Pt(6)
  expected.append(question)
  d.add_paragraph(p.text)
  expected.append(p.text)
  audit.append({'week':week,'source':filename,'index':i,'question':q.text,'number':number,'text':p.text})
path=out/'MN522 PBL Questions and Answers Weeks 1-11.docx'
d.save(path)
actual=[p.text for p in Document(path).paragraphs]
assert actual==expected,'Text differs from original'
(out/'verification.json').write_text(json.dumps({'answer_count':len(audit),'exact_text_match':True,'answers':audit},ensure_ascii=False,indent=2),encoding='utf-8')
print(path)
print('Exact text verified:',len(audit),'answers')

